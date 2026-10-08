"""Typed OECD TiVA normalization and frozen baseline activity weights."""

from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import polars as pl

from green_debt.artifacts import (
    BuildIdentity,
    InputArtifact,
    TableContract,
    verify_manifest,
    write_authoritative_table,
)
from green_debt.tiva import build_activity_weights


SOURCE_COLUMNS = (
    "DATAFLOW",
    "MEASURE",
    "REF_AREA",
    "ACTIVITY",
    "COUNTERPART_AREA",
    "UNIT_MEASURE",
    "FREQ",
    "TIME_PERIOD",
    "OBS_VALUE",
    "UNIT_MULT",
)
SOURCE_SCHEMA = {
    "dataflow": pl.String,
    "measure": pl.String,
    "ref_area": pl.String,
    "activity": pl.String,
    "counterpart_area": pl.String,
    "unit_measure": pl.String,
    "freq": pl.String,
    "year": pl.Int16,
    "obs_value": pl.Float64,
    "unit_mult": pl.Int8,
}
OUTPUT_SCHEMA = {
    "economy_id": pl.String,
    "year": pl.Int16,
    "activity": pl.String,
    "indicator_id": pl.String,
    "value": pl.Float64,
    "unit": pl.String,
    "source_dataflow": pl.String,
    "source_status": pl.String,
}
MAINLV_DATAFLOW = "OECD.STI.PIE:DSD_TIVA_MAINLV@DF_MAINLV(1.1)"
MAINSH_DATAFLOW = "OECD.STI.PIE:DSD_TIVA_MAINSH@DF_MAINSH(1.1)"
EXPECTED_INDICATORS = {
    "dfd_fva_level": ("USD_millions", tuple(range(1995, 2023))),
    "fd_va_level": ("USD_millions", tuple(range(1995, 2023))),
    "exgr_level": ("USD_millions", tuple(range(2000, 2023))),
    "exgr_dva_level": ("USD_millions", tuple(range(2000, 2023))),
    "exgr_fva_level": ("USD_millions", tuple(range(2000, 2023))),
    "imgr_level": ("USD_millions", tuple(range(2000, 2023))),
    "imgr_dva_level": ("USD_millions", tuple(range(2000, 2023))),
    "prod_level": ("USD_millions", tuple(range(2000, 2023))),
    "valu_level": ("USD_millions", tuple(range(2000, 2023))),
    "exgr_dva_share": ("percent", tuple(range(2000, 2023))),
    "exgr_fva_share": ("percent", tuple(range(2000, 2023))),
    "valu_prod_share": ("percent", tuple(range(2000, 2023))),
}


@dataclass(frozen=True)
class TivaSourceSpec:
    path: Path
    indicator_id: str
    measure: str
    unit_measure: str
    unit_mult: int
    period: tuple[int, int]


@dataclass(frozen=True)
class TivaBuildReport:
    source_files: int
    source_rows: int
    retained_rows: int
    excluded_area_or_aggregate_activity_rows: int
    indicators: tuple[str, ...]
    leaf_activities: int
    unmapped_economies: int
    unmapped_activities: int
    duplicate_keys: int
    negative_values_retained: int
    weight_rows: int
    invalid_weight_groups: int
    output_path: str
    output_bytes: int
    weights_path: str
    weights_bytes: int


@dataclass(frozen=True)
class TivaAuditReport:
    rows: int
    indicators: tuple[str, ...]
    economies: int
    leaf_activities: int
    duplicate_keys: int
    unit_mismatches: int
    indicator_periods: dict[str, tuple[int, ...]]
    weight_rows: int
    weight_versions: tuple[str, ...]
    invalid_weight_groups: int
    status: str


def production_tiva_specs(raw_root: Path) -> tuple[TivaSourceSpec, ...]:
    old = raw_root / "oecd_tiva/20260821"
    supplement = raw_root / "oecd_tiva/20260823/data"
    level = old / "data"
    shares = old / "data_shares"
    return (
        TivaSourceSpec(
            supplement / "DFD_FVA.all_areas_all_activities.world.1995-1999.csv",
            "dfd_fva_level",
            "DFD_FVA",
            "USD",
            6,
            (1995, 1999),
        ),
        TivaSourceSpec(
            level / "DFD_FVA.all_areas_all_activities.world.2000-2022.csv",
            "dfd_fva_level",
            "DFD_FVA",
            "USD",
            6,
            (2000, 2022),
        ),
        TivaSourceSpec(
            supplement / "FD_VA.all_areas_all_activities.world.1995-2022.csv",
            "fd_va_level",
            "FD_VA",
            "USD",
            6,
            (1995, 2022),
        ),
        *(
            TivaSourceSpec(
                level / f"{measure}.all_areas_all_activities.world.2000-2022.csv",
                indicator,
                measure,
                "USD",
                6,
                (2000, 2022),
            )
            for measure, indicator in (
                ("EXGR", "exgr_level"),
                ("EXGR_DVA", "exgr_dva_level"),
                ("EXGR_FVA", "exgr_fva_level"),
                ("IMGR", "imgr_level"),
                ("IMGR_DVA", "imgr_dva_level"),
                ("PROD", "prod_level"),
                ("VALU", "valu_level"),
            )
        ),
        TivaSourceSpec(
            shares
            / "EXGR_DVA.all_areas_all_activities.world.PT_EXGR.2000-2022.csv",
            "exgr_dva_share",
            "EXGR_DVA",
            "PT_EXGR",
            0,
            (2000, 2022),
        ),
        TivaSourceSpec(
            shares
            / "EXGR_FVA.all_areas_all_activities.world.PT_EXGR.2000-2022.csv",
            "exgr_fva_share",
            "EXGR_FVA",
            "PT_EXGR",
            0,
            (2000, 2022),
        ),
        TivaSourceSpec(
            shares
            / "VALU_PROD.all_areas_all_activities.world.PT_PROD_VAL.2000-2022.csv",
            "valu_prod_share",
            "VALU_PROD",
            "PT_PROD_VAL",
            0,
            (2000, 2022),
        ),
    )


def normalize_tiva_csv(
    path: Path,
    *,
    expected_measure: str,
    expected_unit_measure: str = "USD",
    expected_unit_mult: int | None = None,
    expected_period: tuple[int, int] | None = None,
) -> pl.DataFrame:
    """Parse every source column explicitly and freeze its SDMX identity."""

    multiplier = (
        expected_unit_mult
        if expected_unit_mult is not None
        else (6 if expected_unit_measure == "USD" else 0)
    )
    frame = pl.read_csv(
        path,
        schema_overrides={
            "DATAFLOW": pl.String,
            "MEASURE": pl.String,
            "REF_AREA": pl.String,
            "ACTIVITY": pl.String,
            "COUNTERPART_AREA": pl.String,
            "UNIT_MEASURE": pl.String,
            "FREQ": pl.String,
            "TIME_PERIOD": pl.Int16,
            "OBS_VALUE": pl.Float64,
            "UNIT_MULT": pl.Int8,
        },
        infer_schema_length=0,
    )
    if tuple(frame.columns) != SOURCE_COLUMNS:
        raise ValueError(
            f"TiVA source columns changed in {path.name}: {tuple(frame.columns)}"
        )
    output = frame.rename(
        {
            "DATAFLOW": "dataflow",
            "MEASURE": "measure",
            "REF_AREA": "ref_area",
            "ACTIVITY": "activity",
            "COUNTERPART_AREA": "counterpart_area",
            "UNIT_MEASURE": "unit_measure",
            "FREQ": "freq",
            "TIME_PERIOD": "year",
            "OBS_VALUE": "obs_value",
            "UNIT_MULT": "unit_mult",
        }
    ).cast(SOURCE_SCHEMA)
    expected_dataflow = (
        MAINLV_DATAFLOW if expected_unit_measure == "USD" else MAINSH_DATAFLOW
    )
    level_errors = output.filter(
        (pl.col("dataflow") != expected_dataflow)
        | (pl.col("measure") != expected_measure)
        | (pl.col("counterpart_area") != "W")
        | (pl.col("unit_measure") != expected_unit_measure)
        | (pl.col("freq") != "A")
        | (pl.col("unit_mult") != multiplier)
    )
    if level_errors.height:
        fields = []
        for field, expected in (
            ("DATAFLOW", expected_dataflow),
            ("MEASURE", expected_measure),
            ("COUNTERPART_AREA", "W"),
            ("UNIT_MEASURE", expected_unit_measure),
            ("FREQ", "A"),
            ("UNIT_MULT", multiplier),
        ):
            source_field = {
                "DATAFLOW": "dataflow",
                "MEASURE": "measure",
                "COUNTERPART_AREA": "counterpart_area",
                "UNIT_MEASURE": "unit_measure",
                "FREQ": "freq",
                "UNIT_MULT": "unit_mult",
            }[field]
            if output.filter(pl.col(source_field) != expected).height:
                fields.append(field)
        raise ValueError(f"TiVA frozen levels changed: {', '.join(fields)}")
    if output.filter(
        pl.col("obs_value").is_null() | ~pl.col("obs_value").is_finite()
    ).height:
        raise ValueError("TiVA OBS_VALUE contains null or nonfinite values")
    if output.filter(
        pl.col("ref_area").is_null()
        | (pl.col("ref_area") == "")
        | pl.col("activity").is_null()
        | (pl.col("activity") == "")
    ).height:
        raise ValueError("TiVA source contains an empty dimension key")
    if expected_period is not None:
        years = tuple(sorted(output.get_column("year").unique().to_list()))
        expected_years = tuple(range(expected_period[0], expected_period[1] + 1))
        if years != expected_years:
            raise ValueError(
                f"TiVA source period differs for {expected_measure}: {years}"
            )
    duplicates = output.group_by(
        "measure", "ref_area", "activity", "counterpart_area", "year"
    ).len().filter(pl.col("len") > 1).height
    if duplicates:
        raise ValueError(f"duplicate TiVA source keys: {duplicates}")
    return output.sort("ref_area", "activity", "year")


def append_tiva_history(
    history: pl.DataFrame, current: pl.DataFrame
) -> pl.DataFrame:
    """Append DFD_FVA history only after proving complete SDMX compatibility."""

    if history.schema != current.schema or tuple(history.columns) != tuple(current.columns):
        raise ValueError("TiVA history and current schemas differ")
    for field in (
        "dataflow",
        "measure",
        "counterpart_area",
        "unit_measure",
        "freq",
        "unit_mult",
    ):
        history_values = set(history.get_column(field).unique().to_list())
        current_values = set(current.get_column(field).unique().to_list())
        if history_values != current_values:
            raise ValueError(f"TiVA history identity differs: {field}")
    for field in ("ref_area", "activity"):
        if set(history.get_column(field).unique().to_list()) != set(
            current.get_column(field).unique().to_list()
        ):
            raise ValueError(f"TiVA history dimension members differ: {field}")
    keys = ["measure", "ref_area", "activity", "counterpart_area", "year"]
    overlap = history.select(keys).join(current.select(keys), on=keys, how="inner")
    if overlap.height:
        raise ValueError(f"TiVA history has overlapping keys: {overlap.height}")
    return pl.concat((history, current)).sort("ref_area", "activity", "year")


def load_activity_registry(path: Path) -> pl.DataFrame:
    """Read the official workbook's 50 leaves and 30 aggregate activities."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Activities-Branches" not in workbook.sheetnames:
            raise ValueError("TiVA structure workbook lacks Activities-Branches")
        sheet = workbook["Activities-Branches"]
        if "Code in" not in str(sheet["C4"].value) or "Economic activity" not in str(
            sheet["E4"].value
        ):
            raise ValueError("TiVA leaf-activity workbook headers changed")
        if "Code in" not in str(sheet["C59"].value) or "Industry aggregate" not in str(
            sheet["E59"].value
        ):
            raise ValueError("TiVA aggregate-activity workbook headers changed")
        rows: list[dict[str, Any]] = []
        for is_leaf, start, end in ((True, 5, 54), (False, 60, 89)):
            for row_number in range(start, end + 1):
                code = str(sheet.cell(row_number, 3).value or "").strip()
                label = str(sheet.cell(row_number, 5).value or "").strip()
                if not code or not label:
                    raise ValueError(
                        f"TiVA activity workbook has an empty row: {row_number}"
                    )
                rows.append(
                    {"activity": code, "activity_label": label, "is_leaf": is_leaf}
                )
    finally:
        workbook.close()
    output = pl.DataFrame(
        rows,
        schema={
            "activity": pl.String,
            "activity_label": pl.String,
            "is_leaf": pl.Boolean,
        },
    ).sort("activity")
    if output.group_by("activity").len().filter(pl.col("len") > 1).height:
        raise ValueError("TiVA activity workbook contains duplicate codes")
    if output.filter(pl.col("is_leaf")).height != 50:
        raise ValueError("TiVA structure workbook must contain exactly 50 leaf activities")
    if output.filter(~pl.col("is_leaf")).height != 30:
        raise ValueError("TiVA structure workbook must contain exactly 30 aggregates")
    return output


def load_leaf_activities(path: Path) -> pl.DataFrame:
    return load_activity_registry(path).filter(pl.col("is_leaf")).drop("is_leaf")


def _map_source(
    frame: pl.DataFrame,
    *,
    indicator_id: str,
    activities: pl.DataFrame,
    economies: pl.DataFrame,
) -> pl.DataFrame:
    activity_join = frame.join(
        activities.select("activity", "is_leaf").with_columns(
            pl.lit(True).alias("_activity_declared")
        ),
        on="activity",
        how="left",
    )
    missing_activities = activity_join.filter(pl.col("_activity_declared").is_null())
    if missing_activities.height:
        codes = missing_activities.get_column("activity").unique().sort().to_list()
        raise ValueError(f"unmapped TiVA activity codes: {codes}")
    mapping = economies.select(
        pl.col("source_code").cast(pl.String).alias("ref_area"),
        pl.col("economy_id").cast(pl.String),
    ).with_columns(pl.lit(True).alias("_economy_declared"))
    if mapping.group_by("ref_area").len().filter(pl.col("len") > 1).height:
        raise ValueError("TiVA economy mapping has duplicate source codes")
    joined = activity_join.join(mapping, on="ref_area", how="left")
    missing_economies = joined.filter(pl.col("_economy_declared").is_null())
    if missing_economies.height:
        codes = missing_economies.get_column("ref_area").unique().sort().to_list()
        raise ValueError(f"unmapped TiVA economy codes: {codes}")
    unit = "USD_millions" if joined.get_column("unit_measure").item(0) == "USD" else "percent"
    return (
        joined.filter(pl.col("is_leaf") & pl.col("economy_id").is_not_null())
        .select(
            "economy_id",
            "year",
            "activity",
            pl.lit(indicator_id).alias("indicator_id"),
            pl.col("obs_value").alias("value"),
            pl.lit(unit).alias("unit"),
            pl.col("dataflow").alias("source_dataflow"),
            pl.lit("reported").alias("source_status"),
        )
        .cast(OUTPUT_SCHEMA)
    )


def _load_contract(path: Path) -> TableContract:
    payload = json.loads(path.read_text(encoding="utf-8"))
    period_value = payload.get("period")
    period = (
        (int(period_value[0]), int(period_value[1]))
        if period_value is not None
        else None
    )
    return TableContract(
        table_id=str(payload["table_id"]),
        schema_version=str(payload["schema_version"]),
        primary_key=tuple(payload["primary_key"]),
        columns={str(key): str(value) for key, value in payload["columns"].items()},
        units={str(key): str(value) for key, value in payload["units"].items()},
        period=period,
        zero_semantics={
            str(key): str(value)
            for key, value in payload.get("zero_semantics", {}).items()
        },
        transformations=tuple(payload.get("transformations", [])),
    )


def _invalid_weight_groups(weights: pl.DataFrame) -> int:
    return (
        weights.group_by("economy_id", "weight_version")
        .agg(
            pl.col("activity_weight").is_null().sum().alias("nulls"),
            pl.col("activity_weight").sum().alias("weight_sum"),
            pl.len().alias("rows"),
        )
        .filter(
            ((pl.col("nulls") > 0) & (pl.col("nulls") < pl.col("rows")))
            | ((pl.col("nulls") == 0) & ((pl.col("weight_sum") - 1.0).abs() > 1e-12))
        )
        .height
    )


def build_tiva_tables(
    *,
    specs: tuple[TivaSourceSpec, ...],
    structure_workbook: Path,
    economies: pl.DataFrame,
    destination: Path,
    contract_path: Path,
    weights_destination: Path,
    weights_contract_path: Path,
    weight_variants: dict[str, tuple[str, ...]],
    equal_variant_activities: tuple[str, ...],
    inputs: tuple[InputArtifact, ...],
    weight_inputs: tuple[InputArtifact, ...],
    build: BuildIdentity,
) -> TivaBuildReport:
    activities = load_activity_registry(structure_workbook)
    normalized_by_indicator: dict[str, list[pl.DataFrame]] = {}
    source_rows = 0
    for spec in specs:
        if spec.indicator_id not in EXPECTED_INDICATORS:
            raise ValueError(f"unregistered TiVA indicator: {spec.indicator_id}")
        frame = normalize_tiva_csv(
            spec.path,
            expected_measure=spec.measure,
            expected_unit_measure=spec.unit_measure,
            expected_unit_mult=spec.unit_mult,
            expected_period=spec.period,
        )
        source_rows += frame.height
        normalized_by_indicator.setdefault(spec.indicator_id, []).append(frame)
    if set(normalized_by_indicator) != set(EXPECTED_INDICATORS):
        raise ValueError("TiVA production indicator set is incomplete")
    source_frames: dict[str, pl.DataFrame] = {}
    for indicator, frames in normalized_by_indicator.items():
        if indicator == "dfd_fva_level":
            if len(frames) != 2:
                raise ValueError("DFD_FVA requires exactly history and current files")
            ordered = sorted(frames, key=lambda item: item.get_column("year").min())
            source_frames[indicator] = append_tiva_history(ordered[0], ordered[1])
        elif len(frames) == 1:
            source_frames[indicator] = frames[0]
        else:
            raise ValueError(f"unexpected multiple TiVA files for {indicator}")
    outputs = [
        _map_source(
            frame,
            indicator_id=indicator,
            activities=activities,
            economies=economies,
        )
        for indicator, frame in sorted(source_frames.items())
    ]
    output = pl.concat(outputs).sort("economy_id", "year", "activity", "indicator_id")
    duplicates = output.group_by(
        "economy_id", "year", "activity", "indicator_id"
    ).len().filter(pl.col("len") > 1).height
    if duplicates:
        raise ValueError(f"duplicate normalized TiVA keys: {duplicates}")
    manifest = write_authoritative_table(
        output, _load_contract(contract_path), destination, inputs, build
    )

    weights = build_activity_weights(
        output,
        variants=weight_variants,
        equal_variant_activities=equal_variant_activities,
        baseline_period=(2000, 2004),
    )
    weights_manifest = write_authoritative_table(
        weights,
        _load_contract(weights_contract_path),
        weights_destination,
        (
            InputArtifact.from_path(destination),
            *weight_inputs,
        ),
        build,
    )
    return TivaBuildReport(
        source_files=len(specs),
        source_rows=source_rows,
        retained_rows=output.height,
        excluded_area_or_aggregate_activity_rows=source_rows - output.height,
        indicators=tuple(sorted(source_frames)),
        leaf_activities=activities.filter(pl.col("is_leaf")).height,
        unmapped_economies=0,
        unmapped_activities=0,
        duplicate_keys=duplicates,
        negative_values_retained=output.filter(pl.col("value") < 0.0).height,
        weight_rows=weights.height,
        invalid_weight_groups=_invalid_weight_groups(weights),
        output_path=str(destination.resolve()),
        output_bytes=manifest.bytes,
        weights_path=str(weights_destination.resolve()),
        weights_bytes=weights_manifest.bytes,
    )


def audit_tiva_tables(
    *, activity_manifest_path: Path, weights_manifest_path: Path
) -> TivaAuditReport:
    activity_manifest = verify_manifest(activity_manifest_path)
    weights_manifest = verify_manifest(weights_manifest_path)
    frame = pl.read_parquet(activity_manifest.destination)
    weights = pl.read_parquet(weights_manifest.destination)
    indicators = tuple(sorted(frame.get_column("indicator_id").unique().to_list()))
    periods = {
        indicator: tuple(
            sorted(
                frame.filter(pl.col("indicator_id") == indicator)
                .get_column("year")
                .unique()
                .to_list()
            )
        )
        for indicator in indicators
    }
    expected_periods = {
        indicator: values[1] for indicator, values in EXPECTED_INDICATORS.items()
    }
    units = {indicator: values[0] for indicator, values in EXPECTED_INDICATORS.items()}
    unit_mismatches = frame.filter(
        pl.struct("indicator_id", "unit").map_elements(
            lambda row: units.get(row["indicator_id"]) != row["unit"],
            return_dtype=pl.Boolean,
        )
    ).height
    duplicates = frame.group_by(
        "economy_id", "year", "activity", "indicator_id"
    ).len().filter(pl.col("len") > 1).height
    versions = tuple(sorted(weights.get_column("weight_version").unique().to_list()))
    expected_versions = tuple(
        sorted(
            (
                "broad_prod_weight",
                "confirmatory_prod_weight",
                "equipment_only_prod_weight",
                "equal_weight",
            )
        )
    )
    invalid_weights = _invalid_weight_groups(weights)
    if (
        indicators != tuple(sorted(EXPECTED_INDICATORS))
        or periods != expected_periods
        or frame.get_column("activity").n_unique() != 50
        or unit_mismatches
        or duplicates
        or versions != expected_versions
        or invalid_weights
    ):
        raise RuntimeError(
            "TiVA audit failed: "
            f"indicators={indicators}, periods={periods}, units={unit_mismatches}, "
            f"activities={frame.get_column('activity').n_unique()}, "
            f"duplicates={duplicates}, weight_versions={versions}, "
            f"invalid_weights={invalid_weights}"
        )
    return TivaAuditReport(
        rows=frame.height,
        indicators=indicators,
        economies=frame.get_column("economy_id").n_unique(),
        leaf_activities=frame.get_column("activity").n_unique(),
        duplicate_keys=duplicates,
        unit_mismatches=unit_mismatches,
        indicator_periods=periods,
        weight_rows=weights.height,
        weight_versions=versions,
        invalid_weight_groups=invalid_weights,
        status="valid",
    )
